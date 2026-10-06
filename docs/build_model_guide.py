from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "docs" / "Privacy_Policy_Analyzer_AI_Model_Guide.docx"
TEAL = "176B57"
INK = "22313A"
MUTED = "5F6D74"
PALE = "EAF3EF"
SAND = "F7F3E9"
LINE = "D8E2DE"


def set_cell_shading(cell, fill: str) -> None:
    properties = cell._tc.get_or_add_tcPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill)
    properties.append(shading)


def set_cell_margins(cell, top=100, start=120, bottom=100, end=120) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    margins = tc_pr.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        tc_pr.append(margins)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = margins.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            margins.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def add_hyperlink(paragraph, text: str, url: str) -> None:
    part = paragraph.part
    relationship_id = part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), relationship_id)
    run = OxmlElement("w:r")
    properties = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), TEAL)
    properties.append(color)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    properties.append(underline)
    run.append(properties)
    value = OxmlElement("w:t")
    value.text = text
    run.append(value)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def set_repeat_table_header(row) -> None:
    properties = row._tr.get_or_add_trPr()
    repeat = OxmlElement("w:tblHeader")
    repeat.set(qn("w:val"), "true")
    properties.append(repeat)


def set_run_font(run, size=None, bold=None, color=INK, italic=None):
    run.font.name = "Arial"
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Arial")
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Arial")
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    run.font.color.rgb = RGBColor.from_string(color)
    return run


def add_body(doc, text, *, color=INK, size=10.5, after=5, bold_lead=None):
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(after)
    paragraph.paragraph_format.line_spacing = 1.12
    if bold_lead and text.startswith(bold_lead):
        set_run_font(paragraph.add_run(bold_lead), size=size, bold=True, color=color)
        set_run_font(paragraph.add_run(text[len(bold_lead):]), size=size, color=color)
    else:
        set_run_font(paragraph.add_run(text), size=size, color=color)
    return paragraph


def add_bullet(doc, text):
    paragraph = doc.add_paragraph(style="List Bullet")
    paragraph.paragraph_format.left_indent = Inches(0.22)
    paragraph.paragraph_format.first_line_indent = Inches(-0.12)
    paragraph.paragraph_format.space_after = Pt(3)
    paragraph.paragraph_format.line_spacing = 1.08
    set_run_font(paragraph.add_run(text), size=10.1, color=INK)
    return paragraph


def add_heading(doc, text, level=1):
    paragraph = doc.add_paragraph(style=f"Heading {level}")
    paragraph.paragraph_format.keep_with_next = True
    paragraph.paragraph_format.space_before = Pt(8 if level == 1 else 5)
    paragraph.paragraph_format.space_after = Pt(4)
    set_run_font(paragraph.add_run(text), size=16 if level == 1 else 11.5, bold=True, color=TEAL if level == 1 else INK)
    return paragraph


def add_table(doc, headers, rows, widths=None, font_size=9.2):
    table = doc.add_table(rows=1, cols=len(headers))
    table.autofit = False
    table.style = "Table Grid"
    set_repeat_table_header(table.rows[0])
    for index, label in enumerate(headers):
        cell = table.rows[0].cells[index]
        if widths:
            cell.width = Inches(widths[index])
        set_cell_shading(cell, TEAL)
        set_cell_margins(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        paragraph = cell.paragraphs[0]
        paragraph.paragraph_format.space_after = Pt(0)
        set_run_font(paragraph.add_run(label), size=font_size, bold=True, color="FFFFFF")
    for row_index, row in enumerate(rows):
        cells = table.add_row().cells
        for index, value in enumerate(row):
            if widths:
                cells[index].width = Inches(widths[index])
            set_cell_margins(cells[index])
            cells[index].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if row_index % 2 == 1:
                set_cell_shading(cells[index], "F5F8F6")
            paragraph = cells[index].paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.paragraph_format.line_spacing = 1.05
            set_run_font(paragraph.add_run(str(value)), size=font_size, color=INK)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)
    return table


def add_callout(doc, title: str, body: str, fill=PALE):
    table = doc.add_table(rows=1, cols=1)
    table.autofit = False
    table.columns[0].width = Inches(6.95)
    cell = table.cell(0, 0)
    set_cell_shading(cell, fill)
    set_cell_margins(cell, top=150, start=190, bottom=150, end=190)
    first = cell.paragraphs[0]
    first.paragraph_format.space_after = Pt(3)
    set_run_font(first.add_run(title), size=10.5, bold=True, color=TEAL)
    second = cell.add_paragraph()
    second.paragraph_format.space_after = Pt(0)
    second.paragraph_format.line_spacing = 1.08
    set_run_font(second.add_run(body), size=9.7, color=INK)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)


def configure_document(doc: Document) -> None:
    section = doc.sections[0]
    section.top_margin = Inches(0.62)
    section.bottom_margin = Inches(0.62)
    section.left_margin = Inches(0.75)
    section.right_margin = Inches(0.75)
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor.from_string(INK)
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    for name in ("Heading 1", "Heading 2"):
        styles[name].font.name = "Arial"
        styles[name]._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
        styles[name]._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    set_run_font(header.add_run("COMP 496  /  PRIVACY POLICY ANALYZER"), size=8, bold=True, color=MUTED)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run_font(footer.add_run("AI MODEL TECHNICAL GUIDE  |  October 2026  |  Page "), size=8, color=MUTED)
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)


def build() -> None:
    doc = Document()
    configure_document(doc)

    eyebrow = doc.add_paragraph()
    eyebrow.paragraph_format.space_before = Pt(11)
    eyebrow.paragraph_format.space_after = Pt(8)
    set_run_font(eyebrow.add_run("MODEL BRIEF  /  VERSION 1.0.0"), size=9, bold=True, color=TEAL)
    title = doc.add_paragraph()
    title.paragraph_format.space_after = Pt(3)
    title.paragraph_format.keep_with_next = True
    set_run_font(title.add_run("Privacy Policy Analyzer"), size=28, bold=True, color=INK)
    subtitle = doc.add_paragraph()
    subtitle.paragraph_format.space_after = Pt(13)
    set_run_font(subtitle.add_run("AI Model Technical Guide"), size=19, bold=True, color=TEAL)
    add_body(doc, "A practical guide to the trained clause tagger, its data, the API handoff, and the limits of the demo build.", color=MUTED, size=11, after=10)
    add_callout(
        doc,
        "What is included",
        "A compact OPP-115-trained topic classifier is integrated into the Flask API and React interface. It suggests policy topics for individual clauses. The existing versioned rules engine remains responsible for findings and the weighted Trust Score.",
    )
    add_heading(doc, "Model at a glance")
    add_table(
        doc,
        ["Task", "Training", "Held-out baseline", "Runtime"],
        [["Multi-label clause topic suggestions", "3,790 segments / 115 policies", "Micro F1 0.7119 at 0.5", "Python standard library; 6.75 MB JSON"]],
        widths=[1.7, 1.7, 1.65, 1.9],
        font_size=9.0,
    )
    add_heading(doc, "Two signals, two jobs")
    add_table(
        doc,
        ["Component", "What it returns", "Effect on Trust Score"],
        [
            ["Clause classifier", "Experimental topic tag(s) attached to policy text", "None"],
            ["Rules and scoring engine", "Evidence-backed findings, five category assessments, and a weighted score", "Calculates the score using the documented weights"],
        ],
        widths=[1.55, 3.15, 2.25],
        font_size=9.1,
    )
    add_body(doc, "The model's topic estimate is not a calibrated probability. It does not establish whether a practice is fair, trustworthy, lawful, or compliant.", color=MUTED, size=9.7, after=0)

    doc.add_page_break()
    add_heading(doc, "1. Training data and method")
    add_body(doc, "The first supervised model uses OPP-115 v1.0, a research corpus of 115 website privacy policies with policy segments annotated by graduate students in law. Consolidated annotations supply the segment labels; policy text and labels are paired by policy ID and segment number. The training pipeline uses 3,790 labeled segments.")
    add_callout(
        doc,
        "Why OPP-115 for this build",
        "The project notes also identify PrivaSeer as a large potential source. PrivaSeer offers broad, mostly unlabeled policy text, while OPP-115 has human topic labels suited to supervised clause classification. This model therefore uses OPP-115; PrivaSeer remains a possible, separately licensed future resource.",
        fill=SAND,
    )
    add_heading(doc, "Training pipeline", level=2)
    for line in (
        "1. Read sanitized policy segments and the 0.5-threshold consolidated annotations from the official corpus archive.",
        "2. Combine annotator labels by segment so a segment can carry more than one topic.",
        "3. Split by whole policy with seed 496: 92 policies for training and 23 unseen policies for validation.",
        "4. Learn one-vs-rest multinomial Naive Bayes token counts with binary unigram and adjacent-bigram features.",
        "5. Measure held-out quality at a fixed 0.5 score threshold, then train the final distributable weights on all 115 policies.",
    ):
        add_body(doc, line, size=9.8, after=3)
    add_heading(doc, "How the research labels map to the UI", level=2)
    add_table(
        doc,
        ["OPP-115 topic label(s)", "Product category"],
        [
            ["First Party Collection/Use", "Data collection"],
            ["Third Party Sharing/Collection", "Data sharing"],
            ["User Choice/Control; User Access, Edit and Deletion; Do Not Track", "User agency"],
            ["Data Retention; Data Security; Policy Change; International and Specific Audiences", "Legal language"],
            ["Other", "No mapped product category"],
            ["Readability", "Not learned from OPP-115; remains a separate text-length heuristic"],
        ],
        widths=[4.4, 2.55],
        font_size=8.8,
    )

    doc.add_page_break()
    add_heading(doc, "2. Evaluation and limits")
    add_body(doc, "The validation set contains 862 segments from 23 policies not used in training. Holding out whole policies reduces leakage from repeated writing within a single policy. Metrics below use the model's fixed 0.5 decision threshold.")
    add_table(
        doc,
        ["Metric", "Held-out result", "How to read it"],
        [
            ["Micro precision", "0.7633", "Of predicted labels, about 76% matched a held-out annotation."],
            ["Micro recall", "0.6669", "The model found about 67% of held-out annotation labels."],
            ["Micro F1", "0.7119", "Combined precision and recall across labels."],
            ["Macro F1", "0.6150", "Average label-level F1; exposes weaker rare categories."],
            ["Exact label-set match", "0.3805", "All labels for a segment matched in about 38% of examples."],
        ],
        widths=[1.55, 1.25, 4.15],
        font_size=8.8,
    )
    add_heading(doc, "Uneven category performance", level=2)
    add_table(
        doc,
        ["Label", "F1", "Positive holdout segments"],
        [
            ["First Party Collection/Use", "0.7671", "356"],
            ["Third Party Sharing/Collection", "0.7654", "255"],
            ["Data Retention", "0.2439", "34"],
            ["User Access, Edit and Deletion", "0.5111", "63"],
            ["Do Not Track", "0.5000", "6"],
        ],
        widths=[3.65, 1.2, 2.1],
        font_size=8.8,
    )
    add_body(doc, "Some categories are substantially weaker. Do Not Track has only six positive validation examples, making its score especially uncertain. Treat this as a baseline for a class demo, not a guarantee for current policies.", color=MUTED, size=9.5)
    add_heading(doc, "Known limitations", level=2)
    for line in (
        "OPP-115 is a small, historical 2016 corpus of English website policies; modern wording and app notices may differ.",
        "The labels describe policy topics, not whether a practice is good, bad, or legally compliant.",
        "Class balance is uneven. The model can miss secondary topics or suggest a topic for ambiguous wording.",
        "Scores are uncalibrated Naive Bayes estimates. No label above the threshold does not mean a clause is unimportant.",
    ):
        add_bullet(doc, line)

    doc.add_page_break()
    add_heading(doc, "3. Integration and demo")
    add_heading(doc, "Request-to-result path", level=2)
    add_table(
        doc,
        ["1  React UI", "2  Flask API", "3  Two analyzers", "4  UI response"],
        [["Policy text is split into clauses", "POST /api/v1/analyses", "Classifier tags clauses; rules engine scores evidence", "Categories, findings, and AI tags are displayed"]],
        widths=[1.7, 1.55, 2.1, 1.6],
        font_size=8.4,
    )
    add_body(doc, "The API response adds `classifier_version` and `clause_classifications` to the existing contract. Each classification contains a clause ID, short excerpt, and up to three labels whose model score is at least 0.5. The interface displays the highest-scoring mapped label for a clause and shows that tags are experimental.")
    add_heading(doc, "Run the local demo", level=2)
    for line in (
        "Terminal 1: python backend/app.py",
        "Terminal 2: pnpm dev",
        "Open the Vite address printed by the UI command. Select Try sample policy or upload public/sample_privacy_policy.txt, then choose Analyze policy.",
        "Point out the Trust Score, rule findings, AI clause tags, and the disclaimer. Explain that topic suggestions do not affect the score.",
    ):
        add_body(doc, line, size=9.6, after=4)
    add_body(doc, "The sample policy is fictional and intentionally mentions collection, advertising sharing, open-ended retention, user requests, and security. The presenter script is in demo/DEMO_SCRIPT.md.", color=MUTED, size=9.5)
    add_heading(doc, "Privacy and licensing", level=2)
    add_body(doc, "The API processes text in memory and does not persist submitted policy text. Logs omit text and excerpts. Keep API authentication and request limits enabled outside local development.")
    add_body(doc, "The OPP-115 project limits use to research, teaching, and scholarship, with terms in the spirit of CC BY-NC. This package is for the team's non-commercial academic demo; it does not grant commercial rights. The original policy archive is not included. Ask the dataset rights holders before commercial use or broader redistribution.", size=9.4)
    sources = doc.add_paragraph()
    sources.paragraph_format.space_before = Pt(2)
    sources.paragraph_format.space_after = Pt(3)
    set_run_font(sources.add_run("Dataset and citation: "), size=9.1, bold=True, color=INK)
    add_hyperlink(sources, "Usable Privacy Policy Project - OPP-115 and terms", "https://www.usableprivacy.org/data/")
    set_run_font(sources.add_run("  |  Related source: "), size=9.1, color=INK)
    add_hyperlink(sources, "PrivaSeer data and code", "https://privaseer.ist.psu.edu/data")
    add_body(doc, "Wilson et al. (2016), “The Creation and Analysis of a Website Privacy Policy Corpus,” ACL 2016, describes the corpus and annotation work.", size=9.1, after=1)
    citation = doc.add_paragraph()
    citation.paragraph_format.space_after = Pt(0)
    add_hyperlink(citation, "ACL Anthology paper P16-1126", "https://aclanthology.org/P16-1126/")

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
